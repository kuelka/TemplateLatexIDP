# -*- coding: utf-8 -*-
"""
Gera a planilha de revisao do D2 pelo autor (passo 4 do corpus RAG).

Entradas (nesta pasta):
  d2-perguntas.csv         os 51 itens, como gerados por gerar-d2.py
  CONFERENCIA-TEXTOS.md    resultado da conferencia de 02/10/2026 (23 itens)

Saida:
  revisao-autor.xlsx       aba Itens (uma linha por item, com colunas de decisao em
                           branco) e aba Decisoes (questoes de desenho em aberto)

Uso:  cd dados/d2-rag && py -3 gerar-revisao.py
Nao altera o d2-perguntas.csv: as decisoes do autor voltam ao gerador depois.
"""

import csv
import hashlib
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).parent
CSV = AQUI / "d2-perguntas.csv"
CONF = AQUI / "CONFERENCIA-TEXTOS.md"

itens = list(csv.DictReader(CSV.open(encoding="utf-8")))
sha = hashlib.sha256(CSV.read_bytes()).hexdigest()


def celulas(linha):
    return [c.strip() for c in linha.strip().strip("|").split("|")]


# Linhas das duas tabelas da conferencia, indexadas pelo id do item
# A secao "Decisoes do autor" registra o que ja foi decidido e aplicado no CSV
ressalvas, conferem, decididos = {}, {}, {}
secao_decisoes = False
for linha in CONF.read_text(encoding="utf-8").splitlines():
    if linha.startswith("## "):
        secao_decisoes = linha.startswith("## Decisões do autor")
    if not re.match(r"\| D2b?-\d{3} \|", linha):
        continue
    c = [x.replace("**", "") for x in celulas(linha)]
    if secao_decisoes:
        decididos[c[0]] = f"Decidido em 02/10/2026 e aplicado no d2-perguntas.csv: {c[1]}. {c[2]}"
        continue
    if len(c) == 4 and c[1] in ("Erro de fato", "Ressalva"):
        ressalvas[c[0]] = (c[1], c[2], c[3])
    else:
        conferem[c[0]] = f"{c[1]} — {c[2]}: {c[3]}"
assert len(ressalvas) == 4 and len(conferem) == 19, (len(ressalvas), len(conferem))

DECISOES_ITEM = ["Manter", "Manter com ajuste", "Reescrever", "Excluir"]
COR = {
    "Erro de fato": "F4CCCC",
    "Ressalva": "FFF2CC",
    "Confere": "D9EAD3",
    "Fora da conferência": "EFEFEF",
}

wb = Workbook()

# --- Aba Instrucoes -------------------------------------------------------
ws = wb.active
ws.title = "Instrucoes"
texto = [
    "Revisão do D2 pelo autor",
    "",
    f"Gerada por dados/d2-rag/gerar-revisao.py a partir de d2-perguntas.csv (SHA-256 {sha}).",
    "Nada aqui altera o dataset: as decisões preenchidas voltam ao gerar-d2.py em um segundo passo.",
    "",
    "Aba Itens: uma linha por item (51). Preencha as colunas em azul:",
    "  decisao_autor — Manter / Manter com ajuste / Reescrever / Excluir;",
    "  pergunta_revisada e gabarito_revisado — só quando mudar o texto;",
    "  observacoes_autor — livre.",
    "",
    "Coluna situacao_conferencia:",
    "  Erro de fato / Ressalva / Confere — resultado da conferência de 02/10/2026 contra o texto legível",
    "  (ver CONFERENCIA-TEXTOS.md);",
    "  Fora da conferência — itens cuja fonte é PDF com camada de texto (Res. CVM 30/2021 e Res. CMN",
    "  4.557/2017), que não entraram naquela conferência. A revisão do autor é a primeira leitura deles",
    "  depois da geração.",
    "",
    "Aba Decisoes: questões de desenho do D2 que são do autor. A coluna escolha_autor fica em branco.",
]
for i, t in enumerate(texto, 1):
    ws.cell(i, 1, t)
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 110

# --- Aba Itens -------------------------------------------------------------
ws = wb.create_sheet("Itens")
cab = [
    "id", "subconjunto", "categoria", "nivel", "norma", "dispositivo",
    "pergunta", "gabarito", "arquivo_fonte",
    "situacao_conferencia", "nota_conferencia", "correcao_proposta",
    "decisao_autor", "pergunta_revisada", "gabarito_revisado", "observacoes_autor",
]
larg = [9, 8, 13, 13, 24, 18, 50, 60, 30, 16, 50, 50, 18, 45, 45, 35]
ws.append(cab)
for it in itens:
    if it["id"] in ressalvas:
        sit, nota, corr = ressalvas[it["id"]]
    elif it["id"] in conferem:
        sit, nota, corr = "Confere", conferem[it["id"]], ""
    else:
        sit, nota, corr = "Fora da conferência", "", ""
    dec, obs = ("Manter com ajuste", decididos[it["id"]]) if it["id"] in decididos else ("", "")
    ws.append([it[k] for k in cab[:9]] + [sit, nota, corr, dec, "", "", obs])

azul = PatternFill("solid", fgColor="DDEBF7")
for j, w in enumerate(larg, 1):
    ws.column_dimensions[ws.cell(1, j).column_letter].width = w
for c in ws[1]:
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    row[9].fill = PatternFill("solid", fgColor=COR[row[9].value])
    for c in row[12:]:
        c.fill = azul
for c in ws[1][12:]:
    c.fill = azul
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type="list", formula1='"' + ",".join(DECISOES_ITEM) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"M2:M{len(itens) + 1}")

# --- Aba Decisoes ---------------------------------------------------------
ws = wb.create_sheet("Decisoes")
ws.append(["id", "questao", "contexto", "opcoes", "escolha_autor", "observacoes_autor"])
decisoes = [
    (
        "DEC-1",
        "Dimensionamento do D2",
        "O alinhamento de 21/08/2026 registra \"30 a 50 perguntas\"; o conjunto tem 51 (39 D2 + 12 D2b).",
        "a) excluir um item (a revisão da aba Itens pode indicar qual); "
        "b) ajustar a redação do texto da dissertação para 51.",
    ),
    (
        "DEC-2",
        "Valor de k no recall@k",
        "Ainda não fixado na metodologia. O D2 mede a recuperação antes do LLM; "
        "o k também limita quantos trechos o agente recebe depois.",
        "Um k único ou uma série de valores reportados lado a lado (decisão do autor).",
    ),
    (
        "DEC-3",
        "Correções propostas pela conferência",
        "4 itens: D2b-048 (erro de data de publicação), D2b-050, D2-032 e D2-036 (ressalvas). "
        "O erro do D2b-048 também está em bibliografia/normas-rag-corpus.md, item 6.",
        "Decidir item a item na aba Itens (coluna decisao_autor).",
    ),
    (
        "DEC-4",
        "Descrição do corpus",
        "O normas-rag-corpus.md e o texto descrevem o corpus como 10 normas; "
        "bibliografia/normas/texto/ tem hoje texto legível para 14 documentos "
        "(inclui IN RFB 1.585/2015, Regras e Regulamento do Tesouro Direto e o ADC 67/2025).",
        "Confirmar a nova composição para atualizar o normas-rag-corpus.md e o texto.",
    ),
    (
        "DEC-5",
        "Validação com o orientador",
        "Pendência 2 do README do D2.",
        "Registrar data e resultado quando ocorrer.",
    ),
]
for d in decisoes:
    ws.append(list(d) + ["", ""])
for j, w in enumerate([8, 30, 60, 55, 30, 35], 1):
    ws.column_dimensions[ws.cell(1, j).column_letter].width = w
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=1):
    row[4].fill = azul
    row[5].fill = azul
ws.freeze_panes = "A2"

wb.save(AQUI / "revisao-autor.xlsx")
cont = {}
for it in itens:
    s = ressalvas.get(it["id"], ("Confere",))[0] if it["id"] in ressalvas or it["id"] in conferem else "Fora da conferência"
    cont[s] = cont.get(s, 0) + 1
print("revisao-autor.xlsx:", len(itens), "itens", cont)
